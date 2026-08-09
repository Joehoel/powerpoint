"""Spike 4 (Python side): time the same invert flow in python-pptx."""
import io
import sys
import time

from pptx import Presentation
from pptx.dml.color import RGBColor

path = sys.argv[1] if len(sys.argv) > 1 else "../tests/fixtures/hagar-presentatie.pptx"
data = open(path, "rb").read()

BG = RGBColor(0x1A, 0x1A, 0x2E)
FG = RGBColor(0xE8, 0xE8, 0xF0)


def run_once() -> int:
    prs = Presentation(io.BytesIO(data))
    for slide in prs.slides:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = BG
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = FG
    buf = io.BytesIO()
    prs.save(buf)
    return buf.tell()


run_once()  # warmup
times = []
for _ in range(10):
    t0 = time.perf_counter()
    run_once()
    times.append((time.perf_counter() - t0) * 1000)
times.sort()
print(f"python-pptx invert+save median: {times[len(times) // 2]:.1f}ms (min {times[0]:.1f}ms)")
