"""Validate spike 3 output with python-pptx: file opens, colors actually changed."""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL

path = sys.argv[1] if len(sys.argv) > 1 else "out/spike3-inverted.pptx"
prs = Presentation(path)

bg_ok = 0
fg_ok = 0
fg_total = 0
for slide in prs.slides:
    fill = slide.background.fill
    if fill.type == MSO_FILL.SOLID and fill.fore_color.rgb == RGBColor(0x1A, 0x1A, 0x2E):
        bg_ok += 1
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    fg_total += 1
                    if run.font.color.rgb == RGBColor(0xE8, 0xE8, 0xF0):
                        fg_ok += 1

print(f"opens OK: {len(prs.slides)} slides")
print(f"backgrounds solid+correct: {bg_ok}/{len(prs.slides)}")
print(f"run colors correct: {fg_ok}/{fg_total}")
assert bg_ok == len(prs.slides) and fg_ok == fg_total, "validation FAILED"
print("VALIDATION PASSED")
