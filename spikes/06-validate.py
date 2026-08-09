"""Spike 6 validation: generate a .docx with python-docx, and afterwards verify
the recolored output produced by the JS WordprocessingML mini-layer.

Usage:
    uv run --with python-docx spikes/06-validate.py generate
    uv run --with python-docx spikes/06-validate.py verify
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

here = Path(__file__).parent / "out"

if sys.argv[1] == "generate":
    doc = Document()
    doc.add_heading("Spike 6", level=1)
    doc.add_paragraph("Hello from python-docx. ")  # trailing space: fidelity probe
    doc.add_paragraph("Tweede alinea met wat tekst.")
    doc.save(here / "spike6-input.docx")
    print("generated out/spike6-input.docx")
else:
    doc = Document(here / "spike6-recolored.docx")
    expected = RGBColor(0xC0, 0x39, 0x2B)
    texts = [p.text for p in doc.paragraphs if p.text]
    runs = [r for p in doc.paragraphs for r in p.runs]
    colored = [r for r in runs if r.font.color and r.font.color.rgb == expected]
    trailing = any(r.text.endswith(" ") for r in runs)
    print(f"opens OK: {len(texts)} non-empty paragraphs: {texts}")
    print(f"runs recolored: {len(colored)}/{len(runs)}")
    print(f"trailing space survived round-trip: {trailing}")
    assert len(colored) == len(runs) and trailing, "validation FAILED"
    print("VALIDATION PASSED")
